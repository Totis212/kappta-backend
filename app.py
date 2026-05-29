from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from database import get_connection, sync_connection
from fpdf import FPDF
import openpyxl
from io import BytesIO
import datetime

app = Flask(__name__)
CORS(app)

@app.after_request
def after_request(response):
    response.headers.add('Access-Control-Allow-Origin', '*')
    response.headers.add('Access-Control-Allow-Headers', 'Content-Type')
    response.headers.add('Access-Control-Allow-Methods', 'GET,POST,PUT,DELETE,OPTIONS')
    return response

@app.route('/')
def home():
    return jsonify({"mensaje": "¡Backend de Kappta funcionando!"})

# ===================== LOGIN Y REGISTRO =====================

@app.route('/api/login', methods=['POST', 'OPTIONS'])
def login():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        username = data.get('username')
        password = data.get('password')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "SELECT Id, Username FROM Usuarios WHERE Username = ? AND Password = ?",
            (username, password)
        )
        row = cursor.fetchone()

        cursor.close()
        conn.close()

        if row:
            return jsonify({
                "mensaje": "Login exitoso",
                "usuario": {"id": row['Id'], "username": row['Username']}
            }), 200
        else:
            return jsonify({"error": "Usuario o contraseña incorrectos"}), 401

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/registro', methods=['POST', 'OPTIONS'])
def registro():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        username = data.get('username')
        password = data.get('password')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "INSERT INTO Usuarios (Username, Password) VALUES (?, ?)",
            (username, password)
        )
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Usuario registrado correctamente"}), 201

    except Exception as e:
        return jsonify({"error": "El usuario ya existe"}), 400

# ===================== CATEGORÍAS =====================

@app.route('/api/categorias', methods=['GET'])
def get_categorias():
    try:
        usuario_id = request.args.get('usuarioId')

        conn = get_connection()
        cursor = conn.cursor()

        if usuario_id:
            cursor.execute(
                """
                SELECT c.Id, c.Nombre, c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
                FROM Categorias c
                LEFT JOIN Gastos g ON c.Id = g.CategoriaId
                WHERE c.UsuarioId = ?
                GROUP BY c.Id, c.Nombre, c.Presupuesto
                ORDER BY c.Nombre
                """,
                (usuario_id,)
            )
        else:
            cursor.execute(
                """
                SELECT c.Id, c.Nombre, c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
                FROM Categorias c
                LEFT JOIN Gastos g ON c.Id = g.CategoriaId
                GROUP BY c.Id, c.Nombre, c.Presupuesto
                ORDER BY c.Nombre
                """
            )

        rows = cursor.fetchall()

        categorias = []
        for row in rows:
            gastado = float(row['Gastado'])
            presupuesto = float(row['Presupuesto'])
            categorias.append({
                "id": row['Id'],
                "nombre": row['Nombre'],
                "presupuesto": presupuesto,
                "gastado": gastado,
                "disponible": presupuesto - gastado
            })

        cursor.close()
        conn.close()

        return jsonify(categorias)

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/categorias', methods=['POST', 'OPTIONS'])
def crear_categoria():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        nombre = data.get('nombre')
        presupuesto = data.get('presupuesto')
        usuario_id = data.get('usuarioId')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "INSERT INTO Categorias (Nombre, Presupuesto, UsuarioId) VALUES (?, ?, ?)",
            (nombre, presupuesto, usuario_id)
        )
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Categoría creada correctamente"}), 201

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/categorias/<int:categoria_id>', methods=['PUT', 'OPTIONS'])
def actualizar_presupuesto(categoria_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        nuevo_presupuesto = data.get('presupuesto')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "UPDATE Categorias SET Presupuesto = ? WHERE Id = ?",
            (nuevo_presupuesto, categoria_id)
        )
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Presupuesto actualizado correctamente"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ===================== GASTOS =====================

@app.route('/api/gastos', methods=['POST', 'OPTIONS'])
def agregar_gasto():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        categoria_id = data.get('categoriaId')
        descripcion = data.get('descripcion')
        monto = data.get('monto')
        usuario_id = data.get('usuarioId')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
            FROM Categorias c
            LEFT JOIN Gastos g ON c.Id = g.CategoriaId
            WHERE c.Id = ?
            GROUP BY c.Presupuesto
            """,
            (categoria_id,)
        )
        row = cursor.fetchone()
        presupuesto = float(row['Presupuesto'])
        gastado = float(row['Gastado'])
        disponible = presupuesto - gastado

        if monto > disponible:
            cursor.close()
            conn.close()
            return jsonify({
                "error": "El gasto de ${:.2f} excede el disponible de ${:.2f} en esta categoria".format(monto, disponible),
                "disponible": disponible
            }), 400

        cursor.execute(
            "INSERT INTO Gastos (CategoriaId, Descripcion, Monto, UsuarioId) VALUES (?, ?, ?, ?)",
            (categoria_id, descripcion, monto, usuario_id)
        )
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Gasto agregado correctamente"}), 201

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/gastos/<int:categoria_id>', methods=['GET'])
def get_gastos_por_categoria(categoria_id):
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "SELECT Id, Descripcion, Monto, Fecha FROM Gastos WHERE CategoriaId = ? ORDER BY Fecha DESC",
            (categoria_id,)
        )
        rows = cursor.fetchall()

        gastos = []
        for row in rows:
            gastos.append({
                "id": row['Id'],
                "descripcion": row['Descripcion'],
                "monto": float(row['Monto']),
                "fecha": row['Fecha']
            })

        cursor.close()
        conn.close()

        return jsonify(gastos)

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/gastos/eliminar/<int:gasto_id>', methods=['DELETE', 'OPTIONS'])
def eliminar_gasto(gasto_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute("DELETE FROM Gastos WHERE Id = ?", (gasto_id,))
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Gasto eliminado correctamente"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/gastos/fecha', methods=['GET'])
def get_gastos_por_fecha():
    try:
        fecha_inicio = request.args.get('inicio')
        fecha_fin = request.args.get('fin')

        conn = get_connection()
        cursor = conn.cursor()

        if fecha_inicio and fecha_fin:
            cursor.execute(
                """
                SELECT c.Nombre, g.Descripcion, g.Monto, g.Fecha
                FROM Gastos g
                JOIN Categorias c ON g.CategoriaId = c.Id
                WHERE g.Fecha BETWEEN ? AND ?
                ORDER BY g.Fecha DESC
                """,
                (fecha_inicio, fecha_fin)
            )
        else:
            cursor.execute(
                """
                SELECT c.Nombre, g.Descripcion, g.Monto, g.Fecha
                FROM Gastos g
                JOIN Categorias c ON g.CategoriaId = c.Id
                ORDER BY g.Fecha DESC
                """
            )

        rows = cursor.fetchall()

        gastos = []
        for row in rows:
            gastos.append({
                "categoria": row['Nombre'],
                "descripcion": row['Descripcion'],
                "monto": float(row['Monto']),
                "fecha": row['Fecha']
            })

        cursor.close()
        conn.close()

        return jsonify(gastos)

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ===================== NOMINA =====================

@app.route('/api/usuarios/<int:usuario_id>/nomina', methods=['PUT', 'OPTIONS'])
def guardar_nomina(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        nomina = data.get('nomina')
        periodo = data.get('periodo')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "UPDATE Usuarios SET Nomina = ?, Periodo = ? WHERE Id = ?",
            (nomina, periodo, usuario_id)
        )
        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Nómina guardada correctamente"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/usuarios/<int:usuario_id>/recomendar', methods=['POST', 'OPTIONS'])
def crear_sobres_recomendados(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        nomina = data.get('nomina')

        recomendaciones = [
            {"nombre": "Alimentacion", "porcentaje": 0.30},
            {"nombre": "Vivienda", "porcentaje": 0.25},
            {"nombre": "Transporte", "porcentaje": 0.10},
            {"nombre": "Salud", "porcentaje": 0.10},
            {"nombre": "Ahorro", "porcentaje": 0.10},
            {"nombre": "Educacion", "porcentaje": 0.05},
            {"nombre": "Entretenimiento", "porcentaje": 0.05},
            {"nombre": "Otros", "porcentaje": 0.05}
        ]

        conn = get_connection()
        cursor = conn.cursor()

        for rec in recomendaciones:
            presupuesto = round(nomina * rec["porcentaje"], 2)
            cursor.execute(
                "INSERT INTO Categorias (Nombre, Presupuesto, UsuarioId) VALUES (?, ?, ?)",
                (rec["nombre"], presupuesto, usuario_id)
            )

        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Sobres creados correctamente"}), 201

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/usuarios/<int:usuario_id>/distribuir', methods=['POST', 'OPTIONS'])
def distribuir_nomina(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        nomina = data.get('nomina')
        categorias_ids = data.get('categoriasIds')

        conn = get_connection()
        cursor = conn.cursor()

        if categorias_ids and len(categorias_ids) > 0:
            porcentaje = round(1.0 / len(categorias_ids), 4)
            for cat_id in categorias_ids:
                presupuesto = round(nomina * porcentaje, 2)
                cursor.execute(
                    "UPDATE Categorias SET Presupuesto = Presupuesto + ?, Seleccionada = 1 WHERE Id = ? AND UsuarioId = ?",
                    (presupuesto, cat_id, usuario_id)
                )
            placeholders = ','.join(['?'] * len(categorias_ids))
            cursor.execute(
                "UPDATE Categorias SET Seleccionada = 0 WHERE UsuarioId = ? AND Id NOT IN ({})".format(placeholders),
                [usuario_id] + categorias_ids
            )
        else:
            cursor.execute(
                "SELECT COUNT(*) FROM Categorias WHERE UsuarioId = ?",
                (usuario_id,)
            )
            count = cursor.fetchone()[0]
            if count > 0:
                porcentaje = round(1.0 / count, 4)
                presupuesto = round(nomina * porcentaje, 2)
                cursor.execute(
                    "UPDATE Categorias SET Presupuesto = Presupuesto + ? WHERE UsuarioId = ?",
                    (presupuesto, usuario_id)
                )

        cursor.execute(
            "UPDATE Usuarios SET Nomina = ? WHERE Id = ?",
            (nomina, usuario_id)
        )

        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Nómina distribuida correctamente"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ===================== AHORRO =====================

@app.route('/api/ahorro/<int:usuario_id>', methods=['GET', 'OPTIONS'])
def get_ahorro(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            "SELECT Id, Meta, Acumulado FROM Ahorro WHERE UsuarioId = ?",
            (usuario_id,)
        )
        row = cursor.fetchone()

        cursor.close()
        conn.close()

        if row:
            return jsonify({
                "id": row['Id'],
                "meta": float(row['Meta']),
                "acumulado": float(row['Acumulado'])
            }), 200
        else:
            return jsonify({"meta": 0, "acumulado": 0}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/ahorro', methods=['POST', 'OPTIONS'])
def guardar_ahorro():
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        data = request.get_json()
        usuario_id = data.get('usuarioId')
        meta = data.get('meta')

        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT Id FROM Ahorro WHERE UsuarioId = ?", (usuario_id,))
        row = cursor.fetchone()

        if row:
            cursor.execute(
                "UPDATE Ahorro SET Meta = ? WHERE UsuarioId = ?",
                (meta, usuario_id)
            )
        else:
            cursor.execute(
                "INSERT INTO Ahorro (UsuarioId, Meta, Acumulado) VALUES (?, ?, 0)",
                (usuario_id, meta)
            )

        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({"mensaje": "Meta de ahorro guardada"}), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/ahorro/mover-sobrante/<int:usuario_id>', methods=['POST', 'OPTIONS'])
def mover_sobrante(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT c.Id, c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
            FROM Categorias c
            LEFT JOIN Gastos g ON c.Id = g.CategoriaId
            WHERE c.UsuarioId = ? AND c.EsAhorro = 0
            GROUP BY c.Id, c.Presupuesto
            """,
            (usuario_id,)
        )
        rows = cursor.fetchall()

        total_sobrante = 0
        for row in rows:
            disponible = float(row['Presupuesto']) - float(row['Gastado'])
            if disponible > 0:
                total_sobrante += disponible
                cursor.execute(
                    "UPDATE Categorias SET Presupuesto = ? WHERE Id = ?",
                    (float(row['Gastado']), row['Id'])
                )

        cursor.execute("SELECT Id FROM Ahorro WHERE UsuarioId = ?", (usuario_id,))
        row_ahorro = cursor.fetchone()

        if row_ahorro:
            cursor.execute(
                "UPDATE Ahorro SET Acumulado = Acumulado + ? WHERE UsuarioId = ?",
                (total_sobrante, usuario_id)
            )
        else:
            cursor.execute(
                "INSERT INTO Ahorro (UsuarioId, Meta, Acumulado) VALUES (?, 0, ?)",
                (usuario_id, total_sobrante)
            )

        conn.commit()
        sync_connection(conn)

        cursor.close()
        conn.close()

        return jsonify({
            "mensaje": "Se movieron ${:.2f} al ahorro".format(total_sobrante),
            "sobrante": total_sobrante
        }), 200

    except Exception as e:
        return jsonify({"error": str(e)}), 500

# ===================== REPORTES =====================

@app.route('/api/reportes/semanal/<int:usuario_id>', methods=['GET', 'OPTIONS'])
def reporte_semanal(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                c.Nombre,
                SUM(g.Monto) as TotalGastado,
                COUNT(g.Id) as NumGastos,
                CAST(strftime('%W', g.Fecha) AS INTEGER) as Semana,
                CAST(strftime('%Y', g.Fecha) AS INTEGER) as Anio
            FROM Gastos g
            JOIN Categorias c ON g.CategoriaId = c.Id
            WHERE c.UsuarioId = ? AND g.UsuarioId = ?
            GROUP BY c.Nombre, strftime('%W', g.Fecha), strftime('%Y', g.Fecha)
            ORDER BY Anio DESC, Semana DESC
        """, (usuario_id, usuario_id))
        
        rows = cursor.fetchall()
        resultado = []
        for row in rows:
            resultado.append({
                'categoria': row[0],
                'totalGastado': float(row[1]),
                'numGastos': row[2],
                'semana': row[3],
                'anio': row[4]
            })
        
        cursor.close()
        conn.close()
        return jsonify(resultado)
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@app.route('/api/reportes/mensual/<int:usuario_id>', methods=['GET', 'OPTIONS'])
def reporte_mensual(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    
    try:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("""
            SELECT 
                c.Nombre,
                c.Presupuesto,
                SUM(g.Monto) as TotalGastado,
                COUNT(g.Id) as NumGastos,
                CAST(strftime('%m', g.Fecha) AS INTEGER) as Mes,
                CAST(strftime('%Y', g.Fecha) AS INTEGER) as Anio
            FROM Gastos g
            JOIN Categorias c ON g.CategoriaId = c.Id
            WHERE c.UsuarioId = ? AND g.UsuarioId = ?
            GROUP BY c.Nombre, c.Presupuesto, strftime('%m', g.Fecha), strftime('%Y', g.Fecha)
            ORDER BY Anio DESC, Mes DESC
        """, (usuario_id, usuario_id))
        
        rows = cursor.fetchall()
        resultado = []
        for row in rows:
            gastado = float(row[2])
            presupuesto = float(row[1])
            resultado.append({
                'categoria': row[0],
                'presupuesto': presupuesto,
                'totalGastado': gastado,
                'numGastos': row[3],
                'mes': row[4],
                'anio': row[5],
                'porcentaje': round((gastado / presupuesto) * 100, 1) if presupuesto > 0 else 0,
                'disponible': presupuesto - gastado
            })
        
        cursor.close()
        conn.close()
        return jsonify(resultado)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

# ===================== EXPORTAR =====================

@app.route('/api/exportar/pdf/<int:usuario_id>', methods=['GET', 'OPTIONS'])
def exportar_pdf(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute("SELECT Username FROM Usuarios WHERE Id = ?", (usuario_id,))
        usuario = cursor.fetchone()

        cursor.execute(
            """
            SELECT c.Nombre, c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
            FROM Categorias c
            LEFT JOIN Gastos g ON c.Id = g.CategoriaId
            WHERE c.UsuarioId = ?
            GROUP BY c.Nombre, c.Presupuesto
            ORDER BY c.Nombre
            """,
            (usuario_id,)
        )
        rows = cursor.fetchall()

        cursor.execute("SELECT Meta, Acumulado FROM Ahorro WHERE UsuarioId = ?", (usuario_id,))
        ahorro = cursor.fetchone()

        cursor.close()
        conn.close()

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font('Helvetica', 'B', 20)
        pdf.cell(0, 15, 'Kappta - Reporte de Gastos', ln=True, align='C')
        pdf.set_font('Helvetica', '', 12)
        pdf.cell(0, 8, 'Usuario: ' + str(usuario[0]), ln=True, align='C')
        pdf.cell(0, 8, 'Fecha: ' + datetime.datetime.now().strftime('%d/%m/%Y'), ln=True, align='C')
        pdf.ln(10)

        pdf.set_font('Helvetica', 'B', 12)
        pdf.set_fill_color(108, 92, 231)
        pdf.set_text_color(255, 255, 255)
        pdf.cell(60, 10, 'Categoria', 1, 0, 'C', True)
        pdf.cell(40, 10, 'Presupuesto', 1, 0, 'C', True)
        pdf.cell(40, 10, 'Gastado', 1, 0, 'C', True)
        pdf.cell(40, 10, 'Disponible', 1, 1, 'C', True)

        pdf.set_font('Helvetica', '', 11)
        pdf.set_text_color(0, 0, 0)
        total_presupuesto = 0
        total_gastado = 0

        for row in rows:
            presupuesto = float(row[1])
            gastado = float(row[2])
            disponible = presupuesto - gastado
            total_presupuesto += presupuesto
            total_gastado += gastado

            if disponible < 0:
                pdf.set_text_color(200, 0, 0)
            else:
                pdf.set_text_color(0, 0, 0)

            nombre = row[0].encode('latin-1', 'replace').decode('latin-1')
            pdf.cell(60, 8, nombre, 1, 0, 'C')
            pdf.cell(40, 8, '${:,.2f}'.format(presupuesto), 1, 0, 'C')
            pdf.cell(40, 8, '${:,.2f}'.format(gastado), 1, 0, 'C')
            pdf.cell(40, 8, '${:,.2f}'.format(disponible), 1, 1, 'C')

        pdf.set_font('Helvetica', 'B', 11)
        pdf.set_text_color(108, 92, 231)
        pdf.cell(60, 10, 'TOTAL', 1, 0, 'C')
        pdf.cell(40, 10, '${:,.2f}'.format(total_presupuesto), 1, 0, 'C')
        pdf.cell(40, 10, '${:,.2f}'.format(total_gastado), 1, 0, 'C')
        pdf.cell(40, 10, '${:,.2f}'.format(total_presupuesto - total_gastado), 1, 1, 'C')

        pdf.ln(8)
        pdf.set_text_color(0, 0, 0)
        pdf.set_font('Helvetica', 'B', 14)
        pdf.cell(0, 10, 'Ahorro', ln=True)
        pdf.set_font('Helvetica', '', 12)
        if ahorro:
            pdf.cell(0, 8, 'Meta: ${:,.2f}'.format(float(ahorro[0])), ln=True)
            pdf.cell(0, 8, 'Acumulado: ${:,.2f}'.format(float(ahorro[1])), ln=True)
        else:
            pdf.cell(0, 8, 'Sin datos de ahorro', ln=True)

        output = BytesIO()
        pdf.output(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='reporte_kappta.pdf'
        )

    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route('/api/exportar/excel/<int:usuario_id>', methods=['GET', 'OPTIONS'])
def exportar_excel(usuario_id):
    if request.method == 'OPTIONS':
        return jsonify({}), 200
    try:
        conn = get_connection()
        cursor = conn.cursor()

        cursor.execute(
            """
            SELECT c.Nombre, c.Presupuesto, COALESCE(SUM(g.Monto), 0) AS Gastado
            FROM Categorias c
            LEFT JOIN Gastos g ON c.Id = g.CategoriaId
            WHERE c.UsuarioId = ?
            GROUP BY c.Nombre, c.Presupuesto
            ORDER BY c.Nombre
            """,
            (usuario_id,)
        )
        rows = cursor.fetchall()

        cursor.close()
        conn.close()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Reporte Kappta"

        headers = ['Categoria', 'Presupuesto', 'Gastado', 'Disponible', '% Usado']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
            cell.fill = openpyxl.styles.PatternFill(start_color='6C5CE7', end_color='6C5CE7', fill_type='solid')
            cell.alignment = openpyxl.styles.Alignment(horizontal='center')

        for i, row in enumerate(rows, 2):
            presupuesto = float(row[1])
            gastado = float(row[2])
            disponible = presupuesto - gastado
            porcentaje = round((gastado / presupuesto) * 100, 1) if presupuesto > 0 else 0

            ws.cell(row=i, column=1, value=row[0])
            ws.cell(row=i, column=2, value=presupuesto)
            ws.cell(row=i, column=3, value=gastado)
            ws.cell(row=i, column=4, value=disponible)
            ws.cell(row=i, column=5, value=porcentaje)

        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 12

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='reporte_kappta.xlsx'
        )

    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)